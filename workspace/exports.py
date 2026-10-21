"""Permission-filtered exports and frozen audit bundles.

An explicitly shared audit package grants its frozen record metadata. Evidence files
also require current evidence access at download time. Restricted source material
raises the package classification. No public media route is used.
"""
import csv
import hashlib
import io
import json
import re
import tempfile
import uuid
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from django.apps import apps
from django.conf import settings
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files import File
from django.db import transaction
from django.http import FileResponse, HttpResponse
from django.utils import timezone

from . import models as m
from .permissions import can_view, has_role, require_view, visible_queryset
from .registry import MODULES
from .services import log_event, record_snapshot, soa_snapshot


def spreadsheet_safe(value):
    """Neutralize formula prefixes even when preceded by tabs or whitespace."""
    text = '' if value is None else str(value)
    stripped = text.lstrip('\ufeff \t\r\n\v\f')
    if stripped.startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r', '\n')):
        return "'" + text
    return text


def _json(value):
    return json.dumps(value, ensure_ascii=False, indent=2, default=str)


def _owner(record):
    return (record.owner.get_full_name() or record.owner.username) if record.owner else 'Unassigned'


def _export_rows(actor, kind, queryset):
    if kind not in MODULES or queryset.model is not MODULES[kind].model:
        raise ValidationError('Choose a valid register to export.')
    records = visible_queryset(actor, queryset)
    if kind == 'soa':
        headers = ['SoA', 'Version', 'State', 'Control', 'Original summary', 'Applicability', 'Justification', 'Implementation', 'Effectiveness', 'Owner', 'Omissions']
        rows = []
        for soa in records:
            # Draft exports are a live working view; approved exports use frozen data.
            if soa.status in soa.FROZEN_STATUSES:
                source_rows = soa.snapshot
            else:
                source_rows, _ = soa_snapshot(actor, soa.scope, soa.framework)
            for row in source_rows:
                if row.get('record_id'):
                    implementation = m.ControlImplementation.objects.filter(pk=row['record_id']).first()
                    if implementation is None or not can_view(actor, implementation):
                        row = {'identifier': row.get('identifier', ''), 'summary': row.get('summary', ''), 'omissions': ['Restricted or unavailable implementation omitted from this derivative export']}
                rows.append([soa.identifier, soa.version, soa.status, row.get('identifier', ''), row.get('summary', ''), row.get('applicability', ''), row.get('justification', ''), row.get('implementation_status', ''), row.get('effectiveness', ''), row.get('owner', ''), '; '.join(row.get('omissions', []))])
        return headers, rows
    headers = ['ID', 'Title', 'Description', 'Status', 'Owner', 'Classification', 'Version', 'Updated']
    extras = {
        'risks': [('scenario', 'Scenario'), ('score', 'Initial score'), ('current_score', 'Current score'), ('scoring_rationale', 'Scoring rationale'), ('review_date', 'Review date')],
        'controls': [('control', 'Control'), ('applicability', 'Applicability'), ('justification', 'Justification'), ('implementation_status', 'Implementation'), ('effectiveness', 'Effectiveness'), ('review_date', 'Review date')],
        'evidence': [('checksum', 'SHA-256'), ('collected_on', 'Collected'), ('expires_on', 'Expires'), ('review_notes', 'Review notes')],
        'documents': [('body', 'Content'), ('change_summary', 'Change summary'), ('review_date', 'Review date')],
        'tasks': [('due_date', 'Due date'), ('priority', 'Priority'), ('required_output', 'Required output'), ('completion_notes', 'Completion notes')],
        'scopes': [('statement', 'Scope statement'), ('boundaries', 'Boundaries'), ('climate_relevance', 'Climate relevance'), ('climate_rationale', 'Climate rationale')],
        'findings': [('condition', 'Observed condition'), ('objective_evidence', 'Objective evidence'), ('severity', 'Severity'), ('due_date', 'Due date')],
        'actions': [('root_cause', 'Root cause'), ('plan', 'Action plan'), ('effectiveness_notes', 'Effectiveness check'), ('closure_rationale', 'Closure rationale')],
        'reviews': [('meeting_date', 'Meeting date'), ('minutes', 'Minutes'), ('decisions', 'Decisions'), ('resource_needs', 'Resource needs')],
    }.get(kind, [])
    headers += [title for _, title in extras]
    rows = []
    for record in records:
        row = [record.identifier, record.title, record.description, record.get_status_display(), _owner(record), record.classification, getattr(record, 'version', ''), timezone.localtime(record.updated_at).isoformat()]
        row += [getattr(record, field, '') for field, _ in extras]
        rows.append(row)
    return headers, rows


def export_table(actor, kind, queryset, format):
    if not actor.is_authenticated:
        raise PermissionDenied('Sign in before exporting records.')
    if format not in {'csv', 'xlsx', 'pdf'}:
        raise ValidationError('Supported export formats are CSV, XLSX and PDF.')
    headers, rows = _export_rows(actor, kind, queryset)
    title = MODULES[kind].title
    filename = f'nava-{kind}-{timezone.localdate().isoformat()}.{format}'
    if format == 'csv':
        output = io.StringIO(newline='')
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows([[spreadsheet_safe(value) for value in row] for row in rows])
        response = HttpResponse('\ufeff' + output.getvalue(), content_type='text/csv; charset=utf-8')
    elif format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        book = Workbook()
        sheet = book.active
        sheet.title = title[:31]
        sheet.append(headers)
        for row in rows:
            # Always write user-supplied values as strings, never formulas.
            sheet.append([spreadsheet_safe(value) for value in row])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='29233D')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(16, len(str(column[0].value)) + 4))
        output = io.BytesIO()
        book.save(output)
        response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
        output = io.BytesIO()
        styles = getSampleStyleSheet()
        styles['BodyText'].fontSize = 9
        styles['BodyText'].leading = 13
        report = SimpleDocTemplate(output, pagesize=A4, title=f'NAVA - {title}', rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
        content = [Paragraph(escape(title), styles['Title']), Paragraph(f'Generated {timezone.localtime().strftime("%d %b %Y %H:%M %Z")} · {len(rows)} rows · Permission-filtered records', styles['BodyText']), Paragraph('Internal preparation and recordkeeping. This report is not a certification decision.', styles['BodyText']), Spacer(1, 18)]
        if not rows:
            content.append(Paragraph('No accessible records match this export.', styles['BodyText']))
        for row in rows:
            heading = f'{row[0]} — {row[3]}' if kind == 'soa' else f'{row[0]} — {row[1]}'
            content.append(Paragraph(escape(str(heading)), styles['Heading2']))
            for header, value in zip(headers, row):
                text = str(value or 'Not recorded')
                # Plain text only; reportlab markup is escaped, including URLs.
                for offset in range(0, len(text), 3500):
                    chunk = escape(text[offset:offset + 3500]).replace('\n', '<br/>')
                    content.append(Paragraph(f'<b>{escape(header)}:</b> {chunk}', styles['BodyText']))
            content.append(Spacer(1, 12))
        report.build(content)
        response = HttpResponse(output.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Cache-Control'] = 'private, no-store'
    response['X-Content-Type-Options'] = 'nosniff'
    log_event(actor, 'register.exported', detail=f'{kind}; format={format}; rows={len(rows)}')
    return response


def _filtered_nested_snapshot(actor, value):
    """Render the same permission-filtered derivative used by record detail pages."""
    from .services import filtered_snapshot
    return filtered_snapshot(actor, value)


def _safe_snapshot(actor, record):
    value = json.loads(_json(record_snapshot(record)))
    value = _filtered_nested_snapshot(actor, value)
    for field in record._meta.fields:
        if field.is_relation and field.related_model and issubclass(field.related_model, m.Record):
            linked = getattr(record, field.name)
            if linked and not can_view(actor, linked):
                value[field.name] = None
                value.setdefault('omitted_fields', []).append(field.name)
    for field in record._meta.many_to_many:
        if issubclass(field.related_model, m.Record):
            linked = getattr(record, field.name).all()
            value[field.name] = list(visible_queryset(actor, linked).values_list('pk', flat=True))
    return value


def _collect_records(actor, pack):
    """Collect approved versions plus scoped operating records, preserving IDs."""
    approved = ['approved', 'published', 'superseded']
    records = [pack.scope]
    for model in [m.InterestedParty, m.Requirement, m.Asset, m.Supplier, m.ControlImplementation, m.Task, m.Audit, m.MetricResult, m.GapAssessment, m.RiskAssessmentVersion, m.SoAVersion, m.DocumentVersion, m.ManagementReview, m.CertificationRecord]:
        qs = visible_queryset(actor, model).filter(scope=pack.scope)
        if issubclass(model, m.VersionedRecord):
            qs = qs.filter(status__in=approved)
        if model is m.SoAVersion:
            qs = qs.filter(framework=pack.framework)
        if model is m.ControlImplementation:
            qs = qs.filter(control__framework=pack.framework)
        records.extend(qs.order_by('pk'))
    risks = [row for row in records if isinstance(row, m.RiskAssessmentVersion)]
    risk_ids = [row.pk for row in risks]
    records.extend(visible_queryset(actor, m.RiskMethodologyVersion).filter(pk__in={row.methodology_id for row in risks}, status__in=approved).order_by('pk'))
    records.extend(visible_queryset(actor, m.TreatmentPlan).filter(risk_id__in=risk_ids, status__in=approved).order_by('pk'))
    records.extend(visible_queryset(actor, m.RiskAcceptance).filter(risk_id__in=risk_ids, status__in=approved).order_by('pk'))
    records.extend(visible_queryset(actor, m.Finding).filter(audit__scope=pack.scope).order_by('pk'))
    records.extend(visible_queryset(actor, m.CorrectiveAction).filter(finding__audit__scope=pack.scope).order_by('pk'))
    seen = set()
    result = []
    for record in records:
        key = (record._meta.label_lower, record.pk)
        if key not in seen:
            result.append(record)
            seen.add(key)
    return result


def _snapshot_classifications(value):
    if isinstance(value, dict):
        if value.get('classification') in {'internal', 'confidential', 'restricted'}:
            yield value['classification']
        for nested in value.values():
            yield from _snapshot_classifications(nested)
    elif isinstance(value, list):
        for nested in value:
            yield from _snapshot_classifications(nested)


@transaction.atomic
def build_audit_pack(actor, pack, expected_revision=None):
    if not has_role(actor, 'Manager'):
        raise PermissionDenied('Only an ISMS manager may generate an audit package.')
    current = m.AuditPack.objects.select_for_update().get(pk=pack.pk)
    require_view(actor, current)
    expected = pack.revision if expected_revision is None else expected_revision
    try:
        matches = int(expected) == current.revision
    except (TypeError, ValueError):
        matches = False
    if not matches:
        raise ValidationError('This package changed while you were editing. Reload it before generating.')
    if current.status != 'draft':
        raise ValidationError('Only a draft package can be generated. Frozen packages retain their original versions.')
    require_view(actor, current.scope)
    if current.scope.status not in ['approved', 'published']:
        raise ValidationError('Approve the selected scope before generating an audit package.')
    if not current.framework_id:
        raise ValidationError('Select the framework edition for this package.')
    if current.period_start and current.period_end and current.period_end < current.period_start:
        raise ValidationError('The package coverage end must not precede its start.')
    evidence_rows = list(current.evidence.all().order_by('pk'))
    for evidence in evidence_rows:
        require_view(actor, evidence)
        if evidence.scope_id != current.scope_id:
            raise ValidationError('All selected evidence must belong to the selected scope version.')
    if sum(item.size_bytes for item in evidence_rows) > 256 * 1024 * 1024:
        raise ValidationError('Select at most 256 MB of evidence per package.')
    records = _collect_records(actor, current)
    ranks = {'internal': 0, 'confidential': 1, 'restricted': 2}
    all_classes = [current.classification] + [row.classification for row in records + evidence_rows]
    current.classification = max(all_classes, key=lambda value: ranks.get(value, 2))
    manifest = {
        'schema_version': 1,
        'package': current.identifier,
        'generated_at': timezone.now().isoformat(),
        'generated_by': actor.get_full_name() or actor.username,
        'scope_id': current.scope_id,
        'scope_version': current.scope.version,
        'framework': {'id': current.framework_id, 'name': current.framework.name, 'edition': current.framework.edition, 'catalog_release': current.framework.catalog_release, 'verification_status': current.framework.verification_status, 'source_url': current.framework.source_url, 'provenance': current.framework.provenance},
        'period_start': current.period_start.isoformat() if current.period_start else None,
        'period_end': current.period_end.isoformat() if current.period_end else None,
        'selection_criteria': current.selection_criteria,
        'selection_rule': 'Approved versions and operating records in the selected scope, plus explicitly selected evidence. Coverage dates are descriptive and do not silently filter linked records.',
        'access_policy': 'An explicit package share grants frozen metadata. Every bundled evidence file also requires current evidence access; external shares require future expiry.',
        'omissions': ['Unapproved versioned records, inaccessible records and links, and unselected evidence files are omitted. Evidence outside the requested coverage period remains explicitly labeled.'],
        'records': [_safe_snapshot(actor, record) for record in records],
        'evidence': [],
    }
    nested_classes = list(_snapshot_classifications(manifest['records']))
    if nested_classes:
        current.classification = max([current.classification, *nested_classes], key=lambda value: ranks[value])
    output = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
    saved_name = None
    try:
        with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            for record, snapshot in zip(records, manifest['records']):
                path = f'records/{record._meta.model_name}/{record.identifier}-v{getattr(record, "version", 1)}.json'
                archive.writestr(path, _json(snapshot).encode('utf-8'))
                snapshot['archive_path'] = path
            for evidence in evidence_rows:
                entry = _safe_snapshot(actor, evidence)
                entry['archive_path'] = None
                if evidence.file:
                    suffix = Path(evidence.original_filename).suffix.lower()
                    suffix = suffix if re.fullmatch(r'\.[a-z0-9]{1,10}', suffix) else '.bin'
                    path = f'evidence/{evidence.identifier}-v{evidence.version}{suffix}'
                    digest = hashlib.sha256()
                    total = 0
                    try:
                        with evidence.file.open('rb') as source, archive.open(path, 'w') as target:
                            while chunk := source.read(1024 * 1024):
                                total += len(chunk)
                                if total > settings.NAVA_MAX_UPLOAD_BYTES:
                                    raise ValidationError('A selected evidence file exceeds the configured upload size limit.')
                                digest.update(chunk)
                                target.write(chunk)
                    except OSError as error:
                        raise ValidationError('A selected evidence file is unavailable. Restore it before generating the package.') from error
                    if digest.hexdigest() != evidence.checksum or total != evidence.size_bytes:
                        raise ValidationError('A selected evidence file failed its stored integrity check. Investigate the original file before generating.')
                    entry['archive_path'] = path
                else:
                    entry['omission'] = 'External or structured evidence; no file bytes were collected by the application.'
                entry['expired_at_generation'] = evidence.is_expired
                manifest['evidence'].append(entry)
            index = io.StringIO(newline='')
            writer = csv.writer(index)
            writer.writerow(['ID', 'Type', 'Title', 'Version', 'Status', 'Archive path', 'SHA-256'])
            for entry in manifest['records'] + manifest['evidence']:
                writer.writerow([spreadsheet_safe(entry.get(key, '')) for key in ['identifier', 'model', 'title', 'version', 'status', 'archive_path', 'checksum']])
            archive.writestr('traceability-index.csv', '\ufeff' + index.getvalue())
            archive.writestr('manifest.json', _json(manifest).encode('utf-8'))
            archive.writestr('README.txt', 'NAVA frozen audit package\n\nRead manifest.json for scope, edition, coverage, exact record versions, omissions and evidence SHA-256 hashes. The traceability index lists readable paths. A checksum detects byte changes; it does not prove authenticity. Evidence review and certification decisions require authorized human judgment.\n\nOnly selected evidence file bytes are included. Record links to unselected evidence remain metadata references. External references were not downloaded or independently verified. This package is a snapshot and later source revisions do not rewrite it.\n')
        output.seek(0)
        digest = hashlib.sha256()
        while chunk := output.read(1024 * 1024):
            digest.update(chunk)
        current.checksum = digest.hexdigest()
        current.manifest = manifest
        current.status = 'complete'
        current.full_clean()
        output.seek(0)
        current.file.save(f'{current.identifier}-{uuid.uuid4().hex}.zip', File(output), save=False)
        saved_name = current.file.name
        current.save()
        log_event(actor, 'audit_pack.generated', current, f'{len(records)} record snapshots and {len(evidence_rows)} selected evidence versions; SHA-256 {current.checksum}.')
    except Exception:
        if saved_name:
            current.file.storage.delete(saved_name)
        raise
    finally:
        output.close()
    pack.refresh_from_db()
    return pack


def pack_download_response(actor, pack):
    require_view(actor, pack)
    if pack.status not in pack.FROZEN_STATUSES or not pack.file:
        raise ValidationError('Generate the package before downloading it.')
    for entry in pack.manifest.get('evidence', []):
        evidence = m.EvidenceVersion.objects.filter(pk=entry.get('id')).first()
        if evidence is None or not can_view(actor, evidence):
            raise PermissionDenied('You need current access to every evidence version included in this package.')
    try:
        source = pack.file.open('rb')
        digest = hashlib.sha256()
        while chunk := source.read(1024 * 1024):
            digest.update(chunk)
        if digest.hexdigest() != pack.checksum:
            source.close()
            raise ValidationError('Package integrity verification failed. Investigate storage before using this artifact.')
        source.seek(0)
    except OSError as error:
        raise ValidationError('The frozen package file is unavailable. Restore the original from backup.') from error
    log_event(actor, 'audit_pack.downloaded', pack)
    response = FileResponse(source, as_attachment=True, filename=f'{pack.identifier}.zip', content_type='application/zip')
    response['Cache-Control'] = 'private, no-store'
    response['X-Content-Type-Options'] = 'nosniff'
    return response
